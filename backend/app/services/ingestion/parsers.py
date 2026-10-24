from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".xlsm",
    ".csv",
    ".tsv",
    ".md",
    ".txt",
}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_BODY_CHARS = 200_000
MAX_TABLE_ROWS = 200
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$", re.M)


@dataclass(slots=True)
class ParsedDocument:
    doc_id: str
    product: str
    section: str
    title: str
    body: str
    content_hash: str
    original_name: str


def parse_upload(
    content: bytes,
    *,
    filename: str,
    product: str,
    section: str,
    title: str = "",
) -> ParsedDocument:
    suffix = validate_upload(filename, len(content))
    validate_file_signature(filename, content)
    if suffix == ".pdf":
        body = _parse_pdf(content, Path(filename).stem)
    elif suffix == ".docx":
        body = _parse_docx(content, Path(filename).stem)
    elif suffix in {".xlsx", ".xlsm"}:
        body = _parse_xlsx(content, Path(filename).stem)
    elif suffix in {".csv", ".tsv"}:
        body = _parse_csv(
            content,
            Path(filename).stem,
            delimiter="\t" if suffix == ".tsv" else ",",
        )
    else:
        body = _decode_text(content)
    body = body.strip()[:MAX_BODY_CHARS]
    if not body:
        raise ValueError("文件未解析出可用正文")
    digest = hashlib.sha256(body.encode("utf-8")).hexdigest()
    identity = hashlib.sha256(
        f"{product}\0{section}\0{digest}".encode("utf-8")
    ).hexdigest()
    return ParsedDocument(
        doc_id=f"upload_{identity[:16]}",
        product=product,
        section=section,
        title=title.strip() or _markdown_title(body) or Path(filename).stem,
        body=body,
        content_hash=digest,
        original_name=Path(filename).name,
    )


def chunk_markdown(body: str, chunk_size: int = 512, overlap: int = 80) -> list[str]:
    if not body.strip():
        return []
    positions = [match.start() for match in _HEADING.finditer(body)]
    blocks: list[str] = []
    if not positions:
        blocks = [body.strip()]
    else:
        prefix = body[: positions[0]].strip()
        if prefix:
            blocks.append(prefix)
        for index, start in enumerate(positions):
            end = positions[index + 1] if index + 1 < len(positions) else len(body)
            block = body[start:end].strip()
            if block:
                blocks.append(block)

    chunks: list[str] = []
    step = max(1, chunk_size - overlap)
    for block in blocks:
        if len(block) <= chunk_size:
            chunks.append(block)
            continue
        for start in range(0, len(block), step):
            piece = block[start : start + chunk_size].strip()
            if piece:
                chunks.append(piece)
            if start + chunk_size >= len(block):
                break
    return chunks


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _markdown_title(body: str) -> str:
    match = _HEADING.search(body)
    return match.group(2).strip() if match else ""


def _rows_to_markdown(rows: list[list[object]]) -> str:
    cleaned = [
        [
            str(cell or "")
            .replace("\n", " ")
            .replace("|", "｜")
            .strip()
            for cell in row
        ]
        for row in rows
    ]
    cleaned = [row for row in cleaned if any(row)]
    if len(cleaned) < 2:
        return ""
    width = max(len(row) for row in cleaned)
    normalized = [row + [""] * (width - len(row)) for row in cleaned]
    header, *body = normalized
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _parse_csv(content: bytes, title: str, *, delimiter: str) -> str:
    reader = csv.reader(io.StringIO(_decode_text(content)), delimiter=delimiter)
    rows = list(reader)
    truncated = len(rows) > MAX_TABLE_ROWS + 1
    table = _rows_to_markdown(rows[: MAX_TABLE_ROWS + 1])
    note = f"\n\n（表格过长，仅保留前 {MAX_TABLE_ROWS} 行）" if truncated else ""
    return f"# {title}\n\n{table}{note}".strip()


def _parse_xlsx(content: bytes, title: str) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    lines = [f"# {title}"]
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))[: MAX_TABLE_ROWS + 1]
        table = _rows_to_markdown([list(row) for row in rows])
        if table:
            lines.extend(["", f"## 工作表：{sheet.title}", "", table])
    workbook.close()
    return "\n".join(lines)


def _parse_docx(content: bytes, title: str) -> str:
    import docx

    document = docx.Document(io.BytesIO(content))
    lines = [f"# {title}"]
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(char for char in style if char.isdigit()) or "2"
            text = "#" * min(int(level) + 1, 6) + " " + text
        lines.append(text)
    for index, table in enumerate(document.tables, start=1):
        markdown = _rows_to_markdown(
            [[cell.text.strip() for cell in row.cells] for row in table.rows]
        )
        if markdown:
            lines.extend(["", f"## 表格 {index}", markdown])
    return "\n".join(lines)


def _parse_pdf(content: bytes, title: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    lines = [f"# {title}"]
    for page_no, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            lines.extend(["", f"## 第 {page_no} 页", text])
    return "\n".join(lines)


def validate_upload(filename: str, size: int) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("不支持的文件类型")
    if size <= 0:
        raise ValueError("文件不能为空")
    if size > MAX_UPLOAD_BYTES:
        raise ValueError("文件不能超过 10 MB")
    return suffix


def validate_file_signature(filename: str, content: bytes) -> None:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("文件内容不是有效的 PDF")
        return
    if suffix in {".docx", ".xlsx", ".xlsm"}:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile as error:
            raise ValueError("文件内容不是有效的 Office 文档") from error
        if "[Content_Types].xml" not in names:
            raise ValueError("Office 文件内容缺少类型清单")
        if suffix == ".docx" and "word/document.xml" not in names:
            raise ValueError("文件内容不是有效的 Word 文档")
        if suffix in {".xlsx", ".xlsm"} and "xl/workbook.xml" not in names:
            raise ValueError("文件内容不是有效的 Excel 文档")
        return
    if b"\x00" in content:
        raise ValueError("文件内容包含二进制控制字符")
    sample = content[:8192]
    control_count = sum(
        byte < 32 and byte not in {9, 10, 13}
        for byte in sample
    )
    if control_count > max(len(sample) * 0.02, 8):
        raise ValueError("文件内容包含过多二进制控制字符")
